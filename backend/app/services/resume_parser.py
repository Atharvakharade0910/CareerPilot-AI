"""Conservative local extraction. Every fact retains an exact source line.
No LLM, inferred seniority, inferred years, or invented facts in Phase 1.
"""

import io, re, zipfile, os
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from app.schemas.profile import Extraction, EvidenceField

SKILLS = {
    "Python": ["python"],
    "FastAPI": ["fastapi"],
    "PostgreSQL": ["postgresql", "postgres"],
    "Scikit-learn": ["scikit-learn", "sklearn"],
    "Generative AI": ["generative ai", "gen ai"],
    "Large Language Models": ["large language models", "llms", "llm"],
    "RAG": ["rag", "retrieval augmented generation", "retrieval-augmented generation"],
    "LangGraph": ["langgraph"],
    "LangChain": ["langchain"],
    "Docker": ["docker"],
    "AWS": ["aws", "amazon web services"],
    "Kubernetes": ["kubernetes"],
    "JavaScript": ["javascript"],
    "TypeScript": ["typescript"],
    "React": ["react"],
    "SQL": ["sql"],
    "Java": ["java"],
    "PyTorch": ["pytorch"],
    "TensorFlow": ["tensorflow"],
    "Machine Learning": ["machine learning"],
    "Git": ["git"],
    "Azure": ["azure"],
    "C++": ["c++"],
    "NLP": ["nlp", "natural language processing"],
}
HEADINGS = {
    "education": "education",
    "academic qualifications": "education",
    "projects": "projects",
    "personal projects": "projects",
    "work experience": "experience",
    "experience": "experience",
    "certifications": "certifications",
    "certificates": "certifications",
    "achievements": "achievements",
    "skills": "skills",
    "technical skills": "skills",
    "summary": "summary",
    "profile": "summary",
}


def document_text(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("This file is not a valid PDF.")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Please upload an unencrypted PDF.")
        if len(reader.pages) > 30:
            raise ValueError("Please use a resume with at most 30 pages.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if len(text.strip()) < 40:
            try:
                import fitz
                import pytesseract
                from PIL import Image, ImageFilter, ImageOps
                if os.environ.get("TESSERACT_CMD"):
                    pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
                doc = fitz.open(stream=data, filetype="pdf")
                pages = []
                for page in doc:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
                    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    image = ImageOps.grayscale(image)
                    image = ImageOps.autocontrast(image).filter(ImageFilter.SHARPEN)
                    pages.append(pytesseract.image_to_string(image, config="--psm 3"))
                text = "\n".join(pages)
            except Exception as exc:
                raise ValueError("This scanned PDF needs local OCR. Install Tesseract OCR and the backend OCR dependencies, then retry.") from exc
    elif suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Expanded document is too large.")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("Invalid Word document.")
        doc = Document(io.BytesIO(data))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for row in t.rows for c in row.cells]
        )
    else:
        raise ValueError("Only PDF and DOCX resumes are supported.")
    if len(text.strip()) < 40:
        raise ValueError(
            "No readable resume text found. Scanned PDFs need OCR; upload a text-based PDF or DOCX."
        )
    if len(text) > 100000:
        raise ValueError("Resume text is too long.")
    return text


def extract(text: str) -> Extraction:
    fields = []
    seen = set()
    section = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = line.lower().strip(": ")
        if heading in HEADINGS:
            section = HEADINGS[heading]
            continue
        for skill, aliases in SKILLS.items():
            if skill not in seen and any(
                re.search(r"(?<!\w)" + re.escape(a) + r"(?!\w)", line, re.I)
                for a in aliases
            ):
                # A mention is evidence of a mention, never a verified proficiency claim.
                if re.search(
                    r"\b(no experience|not familiar|want to learn|learning goals|missing skills)\b",
                    line,
                    re.I,
                ):
                    continue
                seen.add(skill)
                fields.append(
                    EvidenceField(
                        value=skill, category="skills", evidence=line, confidence=0.95
                    )
                )
        if section and section not in ("skills", "summary"):
            fields.append(
                EvidenceField(
                    value=line, category=section, evidence=line, confidence=0.75
                )
            )
        for email in re.findall(r"[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}", line):
            fields.append(
                EvidenceField(
                    value=email, category="contact", evidence=line, confidence=0.99
                )
            )
    return Extraction(
        fields=fields,
        warnings=[
            "Review every extracted item. A skill mention does not establish proficiency or years of experience.",
            "Section-based extraction may split multi-line entries. Scanned PDFs use local high-resolution OCR; review OCR-derived facts carefully.",
        ],
    )


def validate_claim(claim: str, evidence: str, source: str) -> bool:
    return bool(
        claim.strip() and evidence.strip() and evidence in source and claim in evidence
    )
